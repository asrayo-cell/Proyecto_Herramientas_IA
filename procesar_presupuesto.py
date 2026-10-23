import pandas as pd
import openpyxl

# 1. Cargar bases de datos de precios
file_precios = 'PRECIO_MATERIALES.xlsx'
df_mat = pd.read_excel(file_precios, sheet_name='PRECIOS_MATERIALES')
df_gen = pd.read_excel(file_precios, sheet_name='PRECIOS_GENERALES', header=None)

# Diccionario de búsqueda de precios
dict_precios = {}

for idx, row in df_mat.iterrows():
    if idx == 0: 
        continue
    desc = str(row.iloc[0]).strip().lower()
    dict_precios[desc] = row.iloc[2]

for idx, row in df_gen.iterrows():
    desc = str(row.iloc[0]).strip().lower()
    dict_precios[desc] = row.iloc[2]

# 2. Cargar el presupuesto original
wb = openpyxl.load_workbook('ITEMIZADO_PRESUPUESTO_A_COMPLETAR.xlsx')
ws = wb['Formulario Oficial']

# 3. Recorrer y completar partidas (Filas 7 a 382)
for r in range(7, 383):
    desc_cell = ws.cell(row=r, column=2).value
    cant_cell = ws.cell(row=r, column=4).value
    pu_cell = ws.cell(row=r, column=5)
    pt_cell = ws.cell(row=r, column=6)
    nf_cell = ws.cell(row=r, column=7)

    if desc_cell and isinstance(desc_cell, str) and cant_cell is not None:
        key = desc_cell.strip().lower()

        if pu_cell.value is None or str(pu_cell.value).strip() == '':
            if key in dict_precios:
                pu_cell.value = dict_precios[key]
            else:
                nf_cell.value = "SI"

        pt_cell.value = f"=D{r}*E{r}"

# 4. Asignar Fórmulas Financieras en el cuadro final
ws['F384'] = "=SUM(F7:F383)"     # TOTAL COSTO DIRECTO
ws['F385'] = "=F384*E385"       # GASTOS GENERALES (15%)
ws['F386'] = "=F384*E386"       # UTILIDADES (15%)
ws['F387'] = "=F384+F385+F386"  # COSTO TOTAL
ws['F388'] = "=F387*E388"       # IVA (19%)
ws['F389'] = "=F387+F388"       # VALOR TOTAL DEL PROYECTO

# Guardar archivo completado
wb.save('ITEMIZADO_PRESUPUESTO_COMPLETADO.xlsx')
print("¡Presupuesto procesado exitosamente!")